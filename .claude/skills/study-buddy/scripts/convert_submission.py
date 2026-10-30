#!/usr/bin/env python3
"""把论文、作业的Markdown转成提交格式（docx、pdf）。

docx：python-docx + markdown-it-py，保留标题层级、段落、粗斜体、行内代码、列表、表格、
      代码块、引用块、图片（相对路径）。公式（$$...$$）按原文保留为文本。
pdf： 先生成docx，再用Word COM导出，所以依赖已安装的Word。
不自动添加封面、页眉页脚、目录；这些按作业要求另行处理。

用法：python convert_submission.py 稿.md --format docx|pdf [-o 输出路径]
输出已存在时不覆盖，另存为-2、-3。
退出码：0成功，2参数或环境错误。
"""
import argparse
import re
import sys
from pathlib import Path

MAX_PIC_INCHES = 6.0
COMMENT = re.compile(r"<!--.*?-->", re.S)


def fail(msg):
    print(msg, file=sys.stderr)
    sys.exit(2)


def add_inline(par, children, base_dir, state=None):
    from docx.shared import Inches
    bold = ital = 0
    for t in children:
        if t.type == "text":
            r = par.add_run(t.content)
            r.bold, r.italic = bool(bold) or None, bool(ital) or None
        elif t.type == "code_inline":
            r = par.add_run(t.content)
            r.font.name = "Consolas"
        elif t.type == "strong_open":
            bold += 1
        elif t.type == "strong_close":
            bold -= 1
        elif t.type == "em_open":
            ital += 1
        elif t.type == "em_close":
            ital -= 1
        elif t.type in ("softbreak", "hardbreak"):
            par.add_run().add_break()
        elif t.type == "image":
            src = t.attrs.get("src", "").strip("<>")
            p = (base_dir / src).resolve()
            if p.exists():
                from PIL import Image
                with Image.open(p) as im:
                    w_in = im.size[0] / (im.info.get("dpi", (96, 96))[0] or 96)
                par.add_run().add_picture(str(p), width=Inches(min(w_in, MAX_PIC_INCHES)))
                if state is not None:
                    state["images"] += 1
            else:
                par.add_run(f"【图片缺失：{src}】")
                if state is not None:
                    state["missing"].append(src)
        elif t.type == "html_inline":
            continue
        elif t.type in ("link_open", "link_close"):
            continue
        else:
            if t.content:
                par.add_run(t.content)


def convert_docx(md_path, out_path):
    from docx import Document
    from markdown_it import MarkdownIt

    text = COMMENT.sub("", md_path.read_text(encoding="utf-8"))
    tokens = MarkdownIt("commonmark").enable("table").parse(text)
    doc = Document()
    base = md_path.parent
    state = {"images": 0, "missing": [], "headings": 0, "tables": 0}
    lists = []  # 栈：("bullet"|"ordered")
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if t.type == "heading_open":
            level = int(t.tag[1])
            inline = tokens[i + 1]
            h = doc.add_heading(level=min(level, 9))
            add_inline(h, inline.children or [], base, state)
            state["headings"] += 1
            i += 3
            continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            lists.append("bullet" if t.type == "bullet_list_open" else "ordered")
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif t.type == "paragraph_open":
            inline = tokens[i + 1]
            style = None
            if lists:
                depth = min(len(lists), 3)
                base_style = "List Bullet" if lists[-1] == "bullet" else "List Number"
                style = base_style if depth == 1 else f"{base_style} {depth}"
            par = doc.add_paragraph(style=style)
            add_inline(par, inline.children or [], base, state)
            i += 3
            continue
        elif t.type == "blockquote_open":
            # 引用块内的段落用缩进表示
            j = i + 1
            while tokens[j].type != "blockquote_close":
                if tokens[j].type == "inline":
                    par = doc.add_paragraph()
                    par.paragraph_format.left_indent = 720 * 635  # 0.5英寸（EMU）
                    add_inline(par, tokens[j].children or [], base, state)
                j += 1
            i = j + 1
            continue
        elif t.type in ("fence", "code_block"):
            par = doc.add_paragraph()
            r = par.add_run(t.content.rstrip("\n"))
            r.font.name = "Consolas"
        elif t.type == "table_open":
            rows, cur, cell_is_head = [], None, False
            j = i + 1
            while tokens[j].type != "table_close":
                tj = tokens[j]
                if tj.type == "tr_open":
                    cur = []
                elif tj.type == "tr_close":
                    rows.append(cur)
                elif tj.type in ("th_open", "td_open"):
                    cell_is_head = tj.type == "th_open"
                elif tj.type == "inline":
                    cur.append((tokens[j], cell_is_head))
                j += 1
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, (tok, head) in enumerate(row):
                    cell_par = table.cell(ri, ci).paragraphs[0]
                    add_inline(cell_par, tok.children or [], base, state)
                    if head:
                        for r in cell_par.runs:
                            r.bold = True
            state["tables"] += 1
            i = j + 1
            continue
        elif t.type == "hr":
            doc.add_paragraph("─" * 30)
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return state


def docx_to_pdf(docx_path, pdf_path):
    try:
        import win32com.client
    except ImportError:
        fail("缺少pywin32，无法导出PDF，见references/setup.md")
    app = win32com.client.DispatchEx("Word.Application")
    app.Visible = False
    try:
        doc = app.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
        try:
            doc.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
        finally:
            doc.Close(False)
    finally:
        app.Quit()


def free_path(p):
    if not p.exists():
        return p
    n = 2
    while True:
        q = p.with_name(f"{p.stem}-{n}{p.suffix}")
        if not q.exists():
            return q
        n += 1


def main():
    ap = argparse.ArgumentParser(description="Markdown转docx/pdf")
    ap.add_argument("markdown", help="最终提交语言的Markdown稿")
    ap.add_argument("--format", choices=["docx", "pdf"], required=True)
    ap.add_argument("-o", "--output", help="输出路径，默认与稿同名；已存在时另存为-2、-3")
    args = ap.parse_args()

    md = Path(args.markdown).resolve()
    if not md.exists():
        fail(f"文件不存在：{md}")
    try:
        import docx  # noqa: F401
        import markdown_it  # noqa: F401
        import PIL  # noqa: F401
    except ImportError as e:
        fail(f"缺少依赖：{e.name}，见references/setup.md")

    out = free_path(Path(args.output).resolve() if args.output else md.with_suffix("." + args.format))
    if args.format == "docx":
        st = convert_docx(md, out)
    else:
        tmp = free_path(out.with_suffix(".docx"))
        st = convert_docx(md, tmp)
        docx_to_pdf(tmp, out)
        tmp.unlink(missing_ok=True)
    print(f"已生成：{out}（标题{st['headings']}个、表格{st['tables']}个、图片{st['images']}张）")
    if st["missing"]:
        print("以下图片不存在，已用占位文字替代：" + "，".join(st["missing"]), file=sys.stderr)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    try:
        main()
        code = 0
    except SystemExit as e:
        code = e.code
    if code in (None, 0):
        try:
            sys.path.insert(0, str(Path(__file__).resolve().parent))
            from check_update import notify_if_update
            notify_if_update()
        except Exception:
            pass
    sys.exit(code)
